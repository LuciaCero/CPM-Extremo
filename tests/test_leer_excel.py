"""
Tests for the Excel reader.

The reader addresses cells by fixed row positions, so these tests build the
spreadsheet from scratch to pin down that layout. See docs/INPUT_FORMAT.md.
"""

from openpyxl import Workbook

from cpm_extremo import input as entrada


def _escribir_excel(ruta, actividades, dn, de, cn, ce, dependencias, indirectos):
    """
    Write a spreadsheet in the exact layout the reader expects.

    Parameters
    ----------
    dependencias : dict
        Maps an activity to the list of activities it precedes.
    indirectos : tuple
        (fixed cost, daily cost). ``None`` leaves the cell empty.
    """

    wb = Workbook()
    ws = wb.active

    def escribir_fila(fila, etiqueta, valores):
        ws.cell(row=fila, column=1, value=etiqueta)
        for col, valor in enumerate(valores, start=2):
            if valor is not None:
                ws.cell(row=fila, column=col, value=valor)

    # Row 1 is the header, so the reader's row 0 is the spreadsheet's row 2.
    escribir_fila(1, "Duración", actividades)
    escribir_fila(2, "normal", dn)
    escribir_fila(3, "extrema", de)
    # Row 4 is intentionally left empty.
    escribir_fila(5, "Costes", actividades)
    escribir_fila(6, "normal", cn)
    escribir_fila(7, "extrema", ce)
    # Row 8 is intentionally left empty.
    escribir_fila(9, "Costes indirectos", ["A+", "BX"])
    escribir_fila(10, None, list(indirectos))
    # Row 11 is intentionally left empty.
    escribir_fila(12, "Dependencias", actividades)

    for i, origen in enumerate(actividades):
        fila = 13 + i
        ws.cell(row=fila, column=1, value=origen)
        for j, destino in enumerate(actividades):
            if destino in dependencias.get(origen, []):
                ws.cell(row=fila, column=2 + j, value=1)

    wb.save(ruta)
    return ruta


def _excel_de_referencia(tmp_path, indirectos=(None, 150)):
    """Write the reference project used across the test suite."""

    return _escribir_excel(
        tmp_path / "proyecto.xlsx",
        actividades=["A", "B", "C", "D"],
        dn=[6, 8, 5, 7],
        de=[4, 6, 3, 5],
        cn=[600, 900, 400, 700],
        ce=[800, 1300, 600, 1100],
        dependencias={"A": ["B", "C"], "B": ["D"], "C": ["D"]},
        indirectos=indirectos,
    )


def test_lee_duraciones_y_costes(tmp_path):
    """Dn, De, Cn and Ce must land in the right columns."""

    df_actividades, _ = entrada.leer_input(_excel_de_referencia(tmp_path))

    assert list(df_actividades["Actividad"]) == ["A", "B", "C", "D"]
    assert list(df_actividades["Dn"]) == [6, 8, 5, 7]
    assert list(df_actividades["De"]) == [4, 6, 3, 5]
    assert list(df_actividades["Cn"]) == [600.0, 900.0, 400.0, 700.0]
    assert list(df_actividades["Ce"]) == [800.0, 1300.0, 600.0, 1100.0]


def test_lee_los_predecesores_de_la_matriz(tmp_path):
    """A 1 at (row X, column Y) means X precedes Y."""

    df_actividades, _ = entrada.leer_input(_excel_de_referencia(tmp_path))
    predecesores = dict(zip(df_actividades["Actividad"], df_actividades["Predecesor"]))

    assert predecesores["A"] == []
    assert predecesores["B"] == ["A"]
    assert predecesores["C"] == ["A"]
    assert sorted(predecesores["D"]) == ["B", "C"]


def test_los_sucesores_se_derivan_de_los_predecesores(tmp_path):
    """The reader inverts the precedence matrix instead of reading it twice."""

    df_actividades, _ = entrada.leer_input(_excel_de_referencia(tmp_path))
    sucesores = dict(zip(df_actividades["Actividad"], df_actividades["Sucesor"]))

    assert sorted(sucesores["A"]) == ["B", "C"]
    assert sucesores["B"] == ["D"]
    assert sucesores["C"] == ["D"]
    assert sucesores["D"] == []


def test_el_numero_de_actividades_se_deduce_de_la_fila_de_duraciones(tmp_path):
    """Trailing empty columns of the template must be ignored."""

    ruta = _escribir_excel(
        tmp_path / "seis.xlsx",
        actividades=["A", "B", "C", "D", "E", "F"],
        dn=[3, 4, 5, 6, 7, 8],
        de=[2, 3, 4, 5, 6, 7],
        cn=[100, 200, 300, 400, 500, 600],
        ce=[150, 250, 350, 450, 550, 650],
        dependencias={"A": ["B"], "B": ["C"], "C": ["D"], "D": ["E"], "E": ["F"]},
        indirectos=(50, 20),
    )

    df_actividades, _ = entrada.leer_input(ruta)

    assert len(df_actividades) == 6


def test_lee_los_costes_indirectos(tmp_path):
    """Fixed and daily indirect costs come from the same row."""

    ruta = _excel_de_referencia(tmp_path, indirectos=(500, 150))
    _, df_proyecto = entrada.leer_input(ruta)

    assert df_proyecto["C_indirecto"].iloc[0] == 500.0
    assert df_proyecto["C_indirecto_diario"].iloc[0] == 150.0


def test_un_coste_indirecto_vacio_vale_cero(tmp_path):
    """The reference project has no fixed indirect cost."""

    _, df_proyecto = entrada.leer_input(_excel_de_referencia(tmp_path))

    assert df_proyecto["C_indirecto"].iloc[0] == 0.0
    assert df_proyecto["C_indirecto_diario"].iloc[0] == 150.0


def test_el_excel_de_ejemplo_da_el_resultado_conocido(tmp_path):
    """
    A spreadsheet written by hand must produce the same optimum as the shipped
    inputs/ejemplo1.xlsx: 19 days and 5650 EUR.
    """

    from cpm_extremo import grafo as G
    from cpm_extremo.cpm import algoritmo

    df_actividades, df_proyecto = entrada.leer_input(_excel_de_referencia(tmp_path))

    g = G.reducir_grafo(G.crear_grafo(df_actividades))
    _, historial_coste_tiempo, idx_optima = algoritmo.algoritmo_cpm(
        g, df_actividades, df_proyecto
    )

    assert historial_coste_tiempo[idx_optima] == (19, 5650.0)
